import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

def generar_excel_resultado(res, df_lista, dict_libros, tolerancia=0.50):
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        res.to_excel(writer, sheet_name="Conciliacion", index=False)
        df_lista.to_excel(writer, sheet_name="Lista_Normalizada", index=False)

        ws = writer.book["Conciliacion"]
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font

        for col_idx in range(1, ws.max_column + 1):
            letra = get_column_letter(col_idx)
            ws.column_dimensions[letra].width = 20

    buffer.seek(0)
    return buffer
