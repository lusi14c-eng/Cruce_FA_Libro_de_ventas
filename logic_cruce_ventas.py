import io
import re
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule


# ============================================================
# CONFIGURACIÓN GENERAL DE STREAMLIT
# ============================================================

st.set_page_config(
    page_title="Conciliación Lista de Facturación vs Libros de Ventas",
    page_icon="📊",
    layout="wide",
)

TIPOS_VALIDOS = ["FAC", "N/C", "N/D"]


# ============================================================
# FUNCIONES DE NORMALIZACIÓN Y LIMPIEZA
# ============================================================

def normalizar_documento(valor):
    if pd.isna(valor):
        return ""
    s = str(valor).strip().upper()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    s = re.sub(r"[\s\-_./]+", "", s)
    if s.isdigit():
        s = s.lstrip("0") or "0"
    return s


def convertir_numero(valor):
    if pd.isna(valor):
        return 0.0
    if isinstance(valor, (int, float, np.integer, np.floating)):
        return float(valor)

    s = str(valor).strip().replace(" ", "")
    s = re.sub(r"[^\d,\.\-\(\)]", "", s)

    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        partes = s.split(",")
        if len(partes[-1]) in (1, 2):
            s = "".join(partes[:-1]) + "." + partes[-1]
        else:
            s = s.replace(",", "")
    elif s.count(".") > 1:
        partes = s.split(".")
        s = "".join(partes[:-1]) + "." + partes[-1]

    try:
        return float(s)
    except Exception:
        return 0.0


# ============================================================
# LECTURA Y PROCESAMIENTO DE ARCHIVOS
# ============================================================

@st.cache_data(show_spinner=False)
def cargar_listado_facturacion(file_bytes, file_name, hoja="Documentos_CC"):
    xls = pd.ExcelFile(io.BytesIO(file_bytes), engine="openpyxl")
    hoja_uso = hoja if hoja in xls.sheet_names else xls.sheet_names[0]
    
    df_raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=hoja_uso,
        engine="openpyxl"
    )

    df = pd.DataFrame()

    # Col A (0): Número | Col B (1): Tipo | Col C (2): Fecha Doc | Col F (5): Cliente | Col G (6): Nombre
    # Col Q (16): Base (Subtotal Bs) | Col S (18): IVA (Iva Bs) | Col T (19): Exento (Flete Bs)
    df["doc_num"] = df_raw.iloc[:, 0].apply(normalizar_documento)
    df["tipo"] = df_raw.iloc[:, 1].astype(str).str.strip().str.upper()

    df["tipo"] = df["tipo"].replace({
        "FACTURA": "FAC", "FACT": "FAC", "F": "FAC",
        "NC": "N/C", "NOTACREDITO": "N/C",
        "ND": "N/D", "NOTADEBITO": "N/D"
    })

    df["clave"] = df["tipo"] + "|" + df["doc_num"]
    df["fecha_lista"] = pd.to_datetime(df_raw.iloc[:, 2], errors="coerce")
    df["codigo_cliente"] = df_raw.iloc[:, 5].fillna("").astype(str)
    df["cliente"] = df_raw.iloc[:, 6].fillna("").astype(str)

    base = df_raw.iloc[:, 16].apply(convertir_numero)
    iva = df_raw.iloc[:, 18].apply(convertir_numero)
    exento = df_raw.iloc[:, 19].apply(convertir_numero)

    # Invertir signo en N/C para alineación contable
    mask_nc = df["tipo"] == "N/C"
    base.loc[mask_nc] = base.loc[mask_nc].abs() * -1
    iva.loc[mask_nc] = iva.loc[mask_nc].abs() * -1
    exento.loc[mask_nc] = exento.loc[mask_nc].abs() * -1

    df["base_lista"] = base
    df["iva_lista"] = iva
    df["exento_lista"] = exento
    df["total_lista"] = base + iva + exento

    return df[(df["doc_num"] != "") & (df["tipo"].isin(TIPOS_VALIDOS))].copy()


@st.cache_data(show_spinner=False)
def cargar_libro_ventas(file_bytes, file_name):
    xls = pd.ExcelFile(io.BytesIO(file_bytes), engine="openpyxl")
    sheet_name = xls.sheet_names[0]

    # skiprows=5 omite membretes para situar los encabezados en la fila 6
    df_raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=sheet_name,
        skiprows=5,
        engine="openpyxl"
    )

    # Col F (5): FAC | Col I (8): N/D | Col J (9): N/C
    col_fac = df_raw.iloc[:, 5].apply(normalizar_documento)
    col_nd = df_raw.iloc[:, 8].apply(normalizar_documento)
    col_nc = df_raw.iloc[:, 9].apply(normalizar_documento)

    doc_num = col_fac.replace("", np.nan).combine_first(
        col_nd.replace("", np.nan)
    ).combine_first(
        col_nc.replace("", np.nan)
    ).fillna("")

    tipo = np.where(col_nc != "", "N/C",
           np.where(col_nd != "", "N/D",
           np.where(col_fac != "", "FAC", "")))

    df = pd.DataFrame()
    df["doc_num"] = doc_num
    df["tipo"] = tipo
    df["clave"] = df["tipo"] + "|" + df["doc_num"]

    # Col B (1): Fecha | Col C (2): RIF | Col D (3): Cliente
    df["fecha_libro"] = pd.to_datetime(df_raw.iloc[:, 1], errors="coerce")
    df["rif_libro"] = df_raw.iloc[:, 2].fillna("").astype(str)
    df["cliente_libro"] = df_raw.iloc[:, 3].fillna("").astype(str)

    # Col N (13): Num Comprobante Retencion | Col AA (26): IVA Retenido por Comprador
    col_num_comprobante = df_raw.iloc[:, 13].fillna("").astype(str).str.strip()
    col_iva_retenido = df_raw.iloc[:, 26].apply(convertir_numero)

    # Col P (15): Exentas | Col R (17): Base Imponible | Col T (19): Impuesto IVA | Col O (14): Total
    exento = df_raw.iloc[:, 15].apply(convertir_numero)
    base = df_raw.iloc[:, 17].apply(convertir_numero)
    iva = df_raw.iloc[:, 19].apply(convertir_numero)
    total = df_raw.iloc[:, 14].apply(convertir_numero)

    # REGLA DE DETECCIÓN DE RETENCIONES DE IVA
    es_retencion = (base == 0) & ((col_num_comprobante != "") | (col_iva_retenido > 0))

    df["es_retencion"] = es_retencion
    df["num_comprobante"] = col_num_comprobante
    df["iva_retenido"] = col_iva_retenido

    total = np.where(total == 0, base + iva + exento, total)

    # Invertir signo si la N/C está positiva en el libro
    mask_nc = df["tipo"] == "N/C"
    base = np.where(mask_nc, -np.abs(base), base)
    iva = np.where(mask_nc, -np.abs(iva), iva)
    exento = np.where(mask_nc, -np.abs(exento), exento)
    total = np.where(mask_nc, -np.abs(total), total)

    df["base_libro"] = base
    df["iva_libro"] = iva
    df["exento_libro"] = exento
    df["total_libro"] = total
    df["periodo"] = file_name.replace(".xlsx", "").replace(".xlsm", "")

    return df[(df["doc_num"] != "") & (df["tipo"].isin(TIPOS_VALIDOS))].copy()


# ============================================================
# LÓGICA DE CONCILIACIÓN
# ============================================================

def conciliar_datos(df_lista, df_libros, tolerancia=0.50):
    # Separar retenciones puras para no sumarlas a la base imponible de la venta
    df_libros_ventas = df_libros[~df_libros["es_retencion"]].copy()
    df_libros_retenciones = df_libros[df_libros["es_retencion"]].copy()

    df_lista["dup_lista"] = df_lista.groupby("clave")["clave"].transform("count") > 1
    df_libros_ventas["dup_libro"] = df_libros_ventas.groupby("clave")["clave"].transform("count") > 1

    lista_agg = df_lista.groupby("clave", as_index=False).agg(
        tipo=("tipo", "first"),
        doc_num=("doc_num", "first"),
        fecha_lista=("fecha_lista", "min"),
        codigo_cliente=("codigo_cliente", "first"),
        cliente=("cliente", "first"),
        base_lista=("base_lista", "sum"),
        iva_lista=("iva_lista", "sum"),
        exento_lista=("exento_lista", "sum"),
        total_lista=("total_lista", "sum"),
        dup_lista=("dup_lista", "max"),
        ocurrencias_lista=("clave", "size")
    )

    libro_agg = df_libros_ventas.groupby("clave", as_index=False).agg(
        tipo=("tipo", "first"),
        doc_num=("doc_num", "first"),
        fecha_libro=("fecha_libro", "min"),
        cliente_libro=("cliente_libro", "first"),
        base_libro=("base_libro", "sum"),
        iva_libro=("iva_libro", "sum"),
        exento_libro=("exento_libro", "sum"),
        total_libro=("total_libro", "sum"),
        dup_libro=("dup_libro", "max"),
        periodo=("periodo", lambda x: " + ".join(sorted(set(str(v) for v in x if pd.notna(v))))),
        ocurrencias_libro=("clave", "size")
    )

    # Agrupar retenciones por clave para asociar información
    ret_agg = df_libros_retenciones.groupby("clave", as_index=False).agg(
        num_comprobante=("num_comprobante", lambda x: " + ".join(sorted(set(str(v) for v in x if pd.notna(v) and str(v) != "")))),
        iva_retenido=("iva_retenido", "sum"),
        tiene_retencion=("es_retencion", "any")
    )

    res = pd.merge(lista_agg, libro_agg, on="clave", how="outer", suffixes=("_lis", "_lib"))
    res = pd.merge(res, ret_agg, on="clave", how="left")

    res["tipo"] = res["tipo_lis"].fillna(res["tipo_lib"])
    res["documento"] = res["doc_num_lis"].fillna(res["doc_num_lib"])
    res["periodo"] = res["periodo"].fillna("No en libro")
    res["tiene_retencion"] = res["tiene_retencion"].fillna(False)

    for col in ["base_lista", "iva_lista", "exento_lista", "total_lista",
                "base_libro", "iva_libro", "exento_libro", "total_libro", "iva_retenido"]:
        res[col] = res[col].fillna(0.0)

    res["diferencia_exento"] = (res["exento_libro"] - res["exento_lista"]).round(2)
    res["diferencia_base"] = (res["base_libro"] - res["base_lista"]).round(2)
    res["diferencia_iva"] = (res["iva_libro"] - res["iva_lista"]).round(2)
    res["diferencia_total"] = (res["total_libro"] - res["total_lista"]).round(2)

    existe_lista = res["ocurrencias_lista"].notna() & (res["ocurrencias_lista"] > 0)
    existe_libro = res["ocurrencias_libro"].notna() & (res["ocurrencias_libro"] > 0)

    res["estado"] = "OK"
    res.loc[existe_lista & ~existe_libro, "estado"] = "FALTA EN LIBRO"
    res.loc[~existe_lista & existe_libro, "estado"] = "SOLO EN LIBRO"

    ambos = existe_lista & existe_libro
    res.loc[ambos & (res["diferencia_total"].abs() > tolerancia), "estado"] = "DIFERENCIA DE MONTO"
    res.loc[ambos & (res["diferencia_total"].abs() <= tolerancia), "estado"] = "OK"

    # Marcar banderas
    es_dup = res["dup_lista"].fillna(False) | res["dup_libro"].fillna(False)
    res.loc[es_dup, "estado"] = res.loc[es_dup, "estado"] + " / DUPLICADO"

    return res


def formato_fecha(df):
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.strftime("%d/%m/%Y")
    return df


# ============================================================
# REPORTE DE AUDITORÍA - EXCEL
# ============================================================

def preparar_dataframe_reporte(df):
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.strftime("%d/%m/%Y")
    return df


def seleccionar_columnas_auditoria(res):
    columnas = [
        "tipo", "documento", "clave",
        "fecha_lista", "fecha_libro",
        "codigo_cliente", "cliente", "cliente_libro",
        "num_comprobante", "iva_retenido",
        "exento_lista", "exento_libro", "diferencia_exento",
        "base_lista", "base_libro", "diferencia_base",
        "iva_lista", "iva_libro", "diferencia_iva",
        "total_lista", "total_libro", "diferencia_total",
        "periodo", "ocurrencias_lista", "ocurrencias_libro", "estado",
    ]
    columnas_existentes = [c for c in columnas if c in res.columns]
    return res[columnas_existentes].copy()


def crear_resumen_ejecutivo(res, df_lista, dict_libros):
    total_lista = len(df_lista)
    total_libros = sum(len(df) for df in dict_libros.values())
    ok = int(res["estado"].str.startswith("OK").sum())
    faltan = int(res["estado"].str.startswith("FALTA EN LIBRO").sum())
    solo_libro = int(res["estado"].str.startswith("SOLO EN LIBRO").sum())
    diferencias = int(res["estado"].str.startswith("DIFERENCIA DE MONTO").sum())
    duplicados = int(res["estado"].str.contains("DUPLICADO", na=False).sum())

    resumen = pd.DataFrame({
        "INDICADOR": [
            "Documentos en Lista de Facturación",
            "Documentos en Libros de Ventas",
            "Documentos conciliados OK",
            "Documentos que faltan en Libros",
            "Documentos que aparecen solo en Libros",
            "Documentos con diferencia de monto",
            "Documentos duplicados",
        ],
        "CANTIDAD": [
            total_lista,
            total_libros,
            ok,
            faltan,
            solo_libro,
            diferencias,
            duplicados,
        ]
    })
    return resumen


def crear_resumen_por_tipo(res):
    filas = []
    for tipo in ["FAC", "N/C", "N/D"]:
        x = res[res["tipo"] == tipo].copy()
        filas.append({
            "TIPO": tipo,
            "DOCUMENTOS": len(x),
            "OK": int(x["estado"].str.startswith("OK").sum()),
            "FALTAN EN LIBROS": int(x["estado"].str.startswith("FALTA EN LIBRO").sum()),
            "SOLO EN LIBROS": int(x["estado"].str.startswith("SOLO EN LIBRO").sum()),
            "DIFERENCIAS": int(x["estado"].str.startswith("DIFERENCIA DE MONTO").sum()),
            "DUPLICADOS": int(x["estado"].str.contains("DUPLICADO", na=False).sum()),
            "TOTAL LISTA": x["total_lista"].sum() if "total_lista" in x else 0,
            "TOTAL LIBRO": x["total_libro"].sum() if "total_libro" in x else 0,
            "DIFERENCIA": x["diferencia_total"].sum() if "diferencia_total" in x else 0,
        })
    return pd.DataFrame(filas)


def crear_resumen_por_quincena(res):
    x = res[res["periodo"].notna() & (res["periodo"] != "No en libro")].copy()
    if x.empty:
        return pd.DataFrame(
            columns=["PERIODO", "DOCUMENTOS", "OK", "FALTAN", "DIFERENCIAS", "DUPLICADOS", "TOTAL LIBRO"]
        )

    resumen = (
        x.groupby("periodo")
        .agg(
            DOCUMENTOS=("documento", "count"),
            OK=("estado", lambda s: s.str.startswith("OK").sum()),
            FALTAN=("estado", lambda s: s.str.startswith("FALTA EN LIBRO").sum()),
            DIFERENCIAS=("estado", lambda s: s.str.startswith("DIFERENCIA DE MONTO").sum()),
            DUPLICADOS=("estado", lambda s: s.str.contains("DUPLICADO", na=False).sum()),
            TOTAL_LIBRO=("total_libro", "sum"),
        )
        .reset_index()
        .rename(columns={"periodo": "PERIODO"})
    )
    return resumen


def crear_control_documentos(res):
    control = res.copy()
    control["EN_LISTA"] = control["ocurrencias_lista"].fillna(0).gt(0)
    control["EN_LIBRO"] = control["ocurrencias_libro"].fillna(0).gt(0)
    control["DIF_BASE"] = control["diferencia_base"].round(2)
    control["DIF_EXENTO"] = control["diferencia_exento"].round(2)
    control["DIF_IVA"] = control["diferencia_iva"].round(2)
    control["DIF_TOTAL"] = control["diferencia_total"].round(2)

    columnas = [
        "tipo", "documento", "clave",
        "EN_LISTA", "EN_LIBRO",
        "fecha_lista", "fecha_libro",
        "periodo",
        "total_lista", "total_libro",
        "DIF_BASE", "DIF_EXENTO", "DIF_IVA", "DIF_TOTAL",
        "estado",
    ]
    columnas = [c for c in columnas if c in control.columns]
    return control[columnas]


def aplicar_estilo_excel(writer, nombre_hoja, color_encabezado="1F4E78"):
    ws = writer.book[nombre_hoja]

    fill = PatternFill("solid", fgColor=color_encabezado)
    font = Font(bold=True, color="FFFFFF")
    borde = Border(bottom=Side(style="thin", color="FFFFFF"))

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.border = borde
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    for col_idx in range(1, ws.max_column + 1):
        letra = get_column_letter(col_idx)
        max_len = 0
        for fila in range(1, min(ws.max_row, 500) + 1):
            valor = ws.cell(fila, col_idx).value
            if valor is not None:
                max_len = max(max_len, len(str(valor)))
        ws.column_dimensions[letra].width = min(max(max_len + 2, 12), 32)

    ws.row_dimensions[1].height = 32

    palabras_monto = ["base", "iva", "exento", "total", "diferencia", "monto", "retenido"]
    for col_idx in range(1, ws.max_column + 1):
        encabezado = str(ws.cell(1, col_idx).value or "").lower()
        if any(p in encabezado for p in palabras_monto):
            for fila in range(2, ws.max_row + 1):
                ws.cell(fila, col_idx).number_format = '#,##0.00;(#,##0.00);-'

    for col_idx in range(1, ws.max_column + 1):
        encabezado = str(ws.cell(1, col_idx).value or "").lower()
        if "fecha" in encabezado:
            for fila in range(2, ws.max_row + 1):
                ws.cell(fila, col_idx).number_format = "dd/mm/yyyy"


def aplicar_colores_estado(ws):
    encabezados = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    if "estado" not in encabezados:
        return

    col_estado = get_column_letter(encabezados["estado"])
    rango = f"{col_estado}2:{col_estado}{ws.max_row}"

    ws.conditional_formatting.add(
        rango,
        CellIsRule(
            operator="equal",
            formula=['"OK"'],
            fill=PatternFill("solid", fgColor="C6EFCE")
        )
    )


def generar_excel_resultado(res, df_lista, dict_libros, tolerancia=0.50):
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:

        # 1. RESUMEN EJECUTIVO
        resumen = crear_resumen_ejecutivo(res, df_lista, dict_libros)
        resumen.to_excel(writer, sheet_name="00_Resumen_Ejecutivo", index=False, startrow=2)

        ws = writer.book["00_Resumen_Ejecutivo"]
        ws["A1"] = "REPORTE DE CONCILIACIÓN DE DOCUMENTOS FISCALES"
        ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="17365D")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

        ws["A9"] = "CRITERIO DE CONCILIACIÓN"
        ws["A9"].font = Font(bold=True, size=12)
        ws["A10"] = "Los documentos se identifican mediante TIPO + NÚMERO DE DOCUMENTO."
        ws["A11"] = "FAC = Factura | N/C = Nota de Crédito | N/D = Nota de Débito"
        ws["A12"] = "Las diferencias monetarias se calculan como LIBRO - LISTA DE FACTURACIÓN."
        ws["A13"] = f"La tolerancia utilizada fue de {tolerancia:.2f} Bs."

        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 20

        # 2. CONCILIACIÓN GENERAL
        detalle = seleccionar_columnas_auditoria(res)
        detalle = preparar_dataframe_reporte(detalle)
        detalle.to_excel(writer, sheet_name="01_Conciliacion", index=False)

        # 3. FALTANTES
        faltantes = detalle[detalle["estado"].str.startswith("FALTA EN LIBRO")].copy()
        faltantes.to_excel(writer, sheet_name="02_Faltan_en_Libros", index=False)

        # 4. SOLO EN LIBROS
        solo_libro = detalle[detalle["estado"].str.startswith("SOLO EN LIBRO")].copy()
        solo_libro.to_excel(writer, sheet_name="03_Solo_en_Libros", index=False)

        # 5. DIFERENCIAS DE MONTO
        diferencias = detalle[detalle["estado"].str.startswith("DIFERENCIA DE MONTO")].copy()
        if not diferencias.empty:
            diferencias["__abs_diferencia"] = diferencias["diferencia_total"].abs()
            diferencias = diferencias.sort_values("__abs_diferencia", ascending=False).drop(columns="__abs_diferencia")
        diferencias.to_excel(writer, sheet_name="04_Diferencias", index=False)

        # 6. DUPLICADOS
        duplicados = detalle[detalle["estado"].str.contains("DUPLICADO", na=False)].copy()
        duplicados.to_excel(writer, sheet_name="05_Duplicados", index=False)

        # 7. RESUMEN POR TIPO
        resumen_tipo = crear_resumen_por_tipo(res)
        resumen_tipo.to_excel(writer, sheet_name="06_Por_Tipo", index=False)

        # 8. RESUMEN POR QUINCENA / ARCHIVO
        resumen_quincena = crear_resumen_por_quincena(res)
        resumen_quincena.to_excel(writer, sheet_name="07_Por_Quincena", index=False)

        # 9. CONTROL DE DOCUMENTOS
        control = crear_control_documentos(res)
        control.to_excel(writer, sheet_name="08_Control_Documentos", index=False)

        # 10. LISTA NORMALIZADA
        lista_export = preparar_dataframe_reporte(df_lista)
        lista_export.to_excel(writer, sheet_name="09_Lista_Normalizada", index=False)

        # 11. LIBROS NORMALIZADOS
        libros_export = pd.concat(dict_libros.values(), ignore_index=True)
        libros_export = preparar_dataframe_reporte(libros_export)
        libros_export.to_excel(writer, sheet_name="10_Libros_Normalizados", index=False)

        # APLICAR FORMATO
        hojas = [
            "00_Resumen_Ejecutivo",
            "01_Conciliacion",
            "02_Faltan_en_Libros",
            "03_Solo_en_Libros",
            "04_Diferencias",
            "05_Duplicados",
            "06_Por_Tipo",
            "07_Por_Quincena",
            "08_Control_Documentos",
            "09_Lista_Normalizada",
            "10_Libros_Normalizados",
        ]

        for hoja in hojas:
            if hoja not in writer.book.sheetnames:
                continue
            aplicar_estilo_excel(writer, hoja)
            ws = writer.book[hoja]
            if hoja in [
                "01_Conciliacion",
                "02_Faltan_en_Libros",
                "03_Solo_en_Libros",
                "04_Diferencias",
                "05_Duplicados",
                "08_Control_Documentos",
            ]:
                aplicar_colores_estado(ws)

        # COLOR DE HOJAS
        colores = {
            "00_Resumen_Ejecutivo": "17365D",
            "01_Conciliacion": "1F4E78",
            "02_Faltan_en_Libros": "C00000",
            "03_Solo_en_Libros": "FFC000",
            "04_Diferencias": "ED7D31",
            "05_Duplicados": "7030A0",
            "06_Por_Tipo": "4472C4",
            "07_Por_Quincena": "4472C4",
            "08_Control_Documentos": "5B9BD5",
            "09_Lista_Normalizada": "70AD47",
            "10_Libros_Normalizados": "70AD47",
        }

        for hoja, color in colores.items():
            if hoja in writer.book.sheetnames:
                writer.book[hoja].sheet_properties.tabColor = color

    buffer.seek(0)
    return buffer


# ============================================================
# INTERFAZ PRINCIPAL DE STREAMLIT
# ============================================================

st.title("📊 Conciliación Lista de Facturación vs Libros de Ventas")
st.caption("Verificación por TIPO + NÚMERO DE DOCUMENTO (FAC / N/C / N/D)")

with st.expander("ℹ️ Mapeo de columnas configurado", expanded=False):
    st.markdown(
        """
        - **Lista de Facturación (`Documentos_CC`)**:
          - Documento: Columna **A** | Tipo: Columna **B**
          - Base Imponible: Columna **Q** (`Subtotal Bs.`)
          - Impuesto IVA: Columna **S** (`Iva Bs`)
          - Ventas Exentas: Columna **T** (`Flete Bs`)
        - **Libros de Ventas** (`skiprows=5`):
          - Factura (**FAC**): Columna **F**
          - Nota de Débito (**N/D**): Columna **I**
          - Nota de Crédito (**N/C**): Columna **J**
          - Base Imponible: Columna **R**
          - Impuesto IVA: Columna **T**
          - Ventas Exentas: Columna **P**
          - Retenciones de IVA: Columna **N** (N° Comprobante) y Columna **AA** (Monto Retenido)
        """
    )

st.header("1. Cargar Archivos")

col1, col2 = st.columns(2)

with col1:
    f_lista = st.file_uploader(
        "Lista de Facturación (1 archivo Excel)",
        type=["xlsx", "xlsm"],
        key="file_lista"
    )

with col2:
    f_libros = st.file_uploader(
        "Libro(s) de Ventas (1 o varios archivos Excel)",
        type=["xlsx", "xlsm"],
        accept_multiple_files=True,
        key="file_libros"
    )

st.sidebar.header("Configuración")

tolerancia = st.sidebar.number_input(
    "Tolerancia de diferencia monetaria (Bs.)",
    min_value=0.0,
    value=0.50,
    step=0.01,
    help="Diferencias menores o iguales a este monto se consideran OK."
)

if not f_lista or not f_libros:
    st.info("👋 Carga la Lista de Facturación y al menos un Libro de Ventas para activar la ejecución.")
    st.stop()

st.markdown("---")
btn_ejecutar = st.button("🚀 Ejecutar Conciliación", type="primary", use_container_width=True)

if "procesado" not in st.session_state:
    st.session_state.procesado = False

if btn_ejecutar or st.session_state.procesado:
    st.session_state.procesado = True

    with st.spinner("⏳ Cargando y procesando los archivos... Por favor espera unos segundos."):
        df_lista = cargar_listado_facturacion(f_lista.getvalue(), f_lista.name)

        dict_libros = {}
        list_libros = []
        for f in f_libros:
            df_l = cargar_libro_ventas(f.getvalue(), f.name)
            dict_libros[f.name] = df_l
            list_libros.append(df_l)

        df_libros_todos = pd.concat(list_libros, ignore_index=True)

        res_conciliacion = conciliar_datos(df_lista, df_libros_todos, tolerancia)

    # ============================================================
    # VISTA DE RESULTADOS
    # ============================================================
    st.success("✅ Conciliación realizada exitosamente.")

    st.header("2. Métricas Generales")

    m1, m2, m3, m4, m5, m6 = st.columns(6)
    m1.metric("Total Comparados", f"{len(res_conciliacion):,}")
    m2.metric("✅ OK", f"{int(res_conciliacion['estado'].str.startswith('OK').sum()):,}")
    m3.metric("🔴 Faltan en Libro", f"{int(res_conciliacion['estado'].str.startswith('FALTA EN LIBRO').sum()):,}")
    m4.metric("🟡 Solo en Libro", f"{int(res_conciliacion['estado'].str.startswith('SOLO EN LIBRO').sum()):,}")
    m5.metric("🟠 Diferencia Monto", f"{int(res_conciliacion['estado'].str.startswith('DIFERENCIA DE MONTO').sum()):,}")
    m6.metric("⚠️ Duplicados", f"{int(res_conciliacion['estado'].str.contains('DUPLICADO', na=False).sum()):,}")

    st.subheader("Resumen por Tipo de Documento")
    resumen_tipo = []
    for tipo in TIPOS_VALIDOS:
        sub = res_conciliacion[res_conciliacion["tipo"] == tipo]
        resumen_tipo.append({
            "Tipo": tipo,
            "Total": len(sub),
            "OK": int(sub["estado"].str.startswith("OK").sum()),
            "Falta en Libro": int(sub["estado"].str.startswith("FALTA EN LIBRO").sum()),
            "Solo en Libro": int(sub["estado"].str.startswith("SOLO EN LIBRO").sum()),
            "Diferencia Monto": int(sub["estado"].str.startswith("DIFERENCIA DE MONTO").sum()),
            "Duplicados": int(sub["estado"].str.contains("DUPLICADO", na=False).sum()),
        })

    st.dataframe(pd.DataFrame(resumen_tipo), use_container_width=True, hide_index=True)

    # ============================================================
    # DETALLE INTERACTIVO
    # ============================================================
    st.header("3. Detalle de Conciliación")

    c1, c2, c3 = st.columns(3)
    with c1:
        f_tipo = st.multiselect("Filtrar Tipo", TIPOS_VALIDOS, default=TIPOS_VALIDOS)
    with c2:
        estados_disp = sorted(res_conciliacion["estado"].dropna().unique().tolist())
        f_estado = st.multiselect("Filtrar Estado", estados_disp, default=estados_disp)
    with c3:
        periodos_disp = sorted(res_conciliacion["periodo"].dropna().unique().tolist())
        f_periodo = st.multiselect("Filtrar Período / Libro", periodos_disp, default=periodos_disp)

    vista = res_conciliacion[
        res_conciliacion["tipo"].isin(f_tipo) &
        res_conciliacion["estado"].isin(f_estado) &
        res_conciliacion["periodo"].isin(f_periodo)
    ].copy()

    cols_ver = [
        "tipo", "documento", "fecha_lista", "fecha_libro", "cliente", "codigo_cliente",
        "num_comprobante", "iva_retenido",
        "exento_lista", "exento_libro", "diferencia_exento",
        "base_lista", "base_libro", "diferencia_base",
        "iva_lista", "iva_libro", "diferencia_iva",
        "total_lista", "total_libro", "diferencia_total",
        "periodo", "estado"
    ]
    cols_ver = [c for c in cols_ver if c in vista.columns]

    st.dataframe(
        formato_fecha(vista[cols_ver]),
        use_container_width=True,
        hide_index=True,
        height=500
    )

    # ============================================================
    # BOTÓN DE DESCARGA EXCEL
    # ============================================================
    st.header("4. Exportar Resultado")

    excel_data = generar_excel_resultado(
        res_conciliacion,
        df_lista,
        dict_libros,
        tolerancia=tolerancia
    )

    st.download_button(
        label="📥 Descargar Reporte de Conciliación Completo en Excel",
        data=excel_data,
        file_name="Conciliacion_Lista_Facturacion_vs_Libro_Ventas.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )
