import io
import re
from datetime import datetime

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st


# ============================================================
# FUNCIONES AUXILIARES DE LIMPIEZA Y PROCESAMIENTO
# ============================================================
def clean_str(val):
    if pd.isna(val):
        return ""
    val_str = str(val).strip()
    if val_str.endswith(".0"):
        val_str = val_str[:-2]
    return val_str


def build_key(row, key_cols):
    parts = []
    for c in key_cols:
        val = clean_str(row.get(c, ""))
        parts.append(val)
    return " | ".join(parts)


def parse_period(val):
    if pd.isna(val):
        return None
    val_str = str(val).strip()
    if not val_str:
        return None

    # Formato AAAAMM (ej. 202401)
    if re.match(r"^\d{6}$", val_str):
        try:
            return pd.to_datetime(val_str, format="%Y%m")
        except Exception:
            pass

    # Formato AAAA-MM o AAAA/MM
    if re.match(r"^\d{4}[-/]\d{1,2}$", val_str):
        try:
            return pd.to_datetime(val_str)
        except Exception:
            pass

    # Intento general
    try:
        return pd.to_datetime(val_str)
    except Exception:
        return None


def calculate_variation_metrics(df_grouped):
    """Calcula variaciones respecto al mes anterior y respecto a la media histórica."""
    df_grouped = df_grouped.sort_values(
        ["empresa", "vista", "segmento", "cuenta", "mes"]
    ).reset_index(drop=True)

    # Identificador único de serie de tiempo
    df_grouped["series_id"] = (
        df_grouped["empresa"].astype(str)
        + "||"
        + df_grouped["vista"].astype(str)
        + "||"
        + df_grouped["segmento"].astype(str)
        + "||"
        + df_grouped["cuenta"].astype(str)
    )

    # Variación vs Mes Anterior
    df_grouped["monto_anterior"] = df_grouped.groupby("series_id")["monto"].shift(1)
    df_grouped["var_abs_mes_anterior"] = (
        df_grouped["monto"] - df_grouped["monto_anterior"]
    )
    df_grouped["var_pct_mes_anterior"] = np.where(
        df_grouped["monto_anterior"] != 0,
        df_grouped["var_abs_mes_anterior"] / df_grouped["monto_anterior"].abs(),
        np.nan,
    )

    # Promedio histórico acumulado (excluyendo el mes actual)
    df_grouped["media_historica"] = df_grouped.groupby("series_id")["monto"].transform(
        lambda x: x.shift(1).expanding().mean()
    )
    df_grouped["var_abs_vs_media"] = df_grouped["monto"] - df_grouped["media_historica"]
    df_grouped["var_pct_vs_media"] = np.where(
        df_grouped["media_historica"] != 0,
        df_grouped["var_abs_vs_media"] / df_grouped["media_historica"].abs(),
        np.nan,
    )

    # Impacto absoluto para ordenamiento
    df_grouped["impacto_abs"] = df_grouped["var_abs_mes_anterior"].abs()

    # Asignación de Alertas
    def assign_alert(row):
        pct = abs(row["var_pct_mes_anterior"]) if pd.notna(row["var_pct_mes_anterior"]) else 0
        var_abs = abs(row["var_abs_mes_anterior"]) if pd.notna(row["var_abs_mes_anterior"]) else 0

        if pct >= 0.50 and var_abs >= 1000:
            return "🔴 CRÍTICA"
        elif pct >= 0.20 and var_abs >= 300:
            return "🟠 IMPORTANTE"
        elif pct >= 0.10:
            return "🟡 MODERADA"
        return "🟢 NORMAL"

    df_grouped["nivel_alerta"] = df_grouped.apply(assign_alert, axis=1)
    return df_grouped


def detail_matches_alert(alert_row, df_detail):
    """Filtra el detalle original que compone una línea alertada."""
    if df_detail is None or df_detail.empty:
        return pd.DataFrame()

    cond = (
        (df_detail["empresa"] == alert_row["empresa"])
        & (df_detail["vista"] == alert_row["vista"])
        & (df_detail["segmento"] == alert_row["segmento"])
        & (df_detail["cuenta"] == alert_row["cuenta"])
        & (df_detail["mes"] == alert_row["mes"])
    )
    return df_detail[cond]


# ============================================================
# MÓDULO PRINCIPAL DE STREAMLIT
# ============================================================
def modulo_analizador_gastos():
    st.title("📊 Analizador de Variaciones de Gastos (por Casas)")
    st.markdown(
        "Herramienta analítica para el control de variaciones mensuales de gastos en **Febeca, Sillaca y Beval**, contemplando vistas **Con / Sin Logística**."
    )

    # Carga de datos
    st.sidebar.header("📁 Carga de Datos")
    uploaded_file = st.sidebar.file_uploader(
        "Cargar archivo de Gastos (Excel/CSV)", type=["xlsx", "xls", "csv"]
    )

    if uploaded_file is None:
        st.info("👈 Por favor, sube un archivo de datos en el panel izquierdo para comenzar.")
        return

    # Lectura del archivo
    try:
        if uploaded_file.name.endswith(".csv"):
            df_raw = pd.read_csv(uploaded_file)
        else:
            df_raw = pd.read_excel(uploaded_file)
    except Exception as e:
        st.error(f"Error al leer el archivo: {e}")
        return

    st.sidebar.success(f"Archivo cargado correctamente ({len(df_raw)} filas).")

    # Mapeo de Columnas
    st.sidebar.subheader("🛠️ Mapeo de Columnas")
    cols = list(df_raw.columns)

    col_empresa = st.sidebar.selectbox("Empresa / Casa", cols, index=0)
    col_vista = st.sidebar.selectbox("Vista (Logística / Sin Logística)", cols, index=min(1, len(cols)-1))
    col_segmento = st.sidebar.selectbox("Segmento / Grupo de Gasto", cols, index=min(2, len(cols)-1))
    col_cuenta = st.sidebar.selectbox("Cuenta / Concepto", cols, index=min(3, len(cols)-1))
    col_mes = st.sidebar.selectbox("Período / Mes", cols, index=min(4, len(cols)-1))
    col_monto = st.sidebar.selectbox("Monto ($ / Bs)", cols, index=min(5, len(cols)-1))

    # Procesamiento y estandarización
    df = df_raw.copy()
    df["empresa"] = df[col_empresa].apply(clean_str)
    df["vista"] = df[col_vista].apply(clean_str)
    df["segmento"] = df[col_segmento].apply(clean_str)
    df["cuenta"] = df[col_cuenta].apply(clean_str)
    df["mes"] = df[col_mes].apply(parse_period)
    df["monto"] = pd.to_numeric(df[col_monto], errors="coerce").fillna(0)

    # Filtrar registros incompletos
    df = df.dropna(subset=["mes"])

    if df.empty:
        st.error("No se pudieron parsear las fechas del archivo. Revisa la columna de 'Mes/Período'.")
        return

    # Agrupación base para métricas de variación
    df_grouped = (
        df.groupby(["empresa", "vista", "segmento", "cuenta", "mes"], as_index=False)["monto"]
        .sum()
    )

    df_analyzed = calculate_variation_metrics(df_grouped)

    # Diagnostics
    diagnostics = pd.DataFrame({
        "Métrica": ["Registros Procesados", "Casas Evaluadas", "Cuentas Analizadas", "Meses Registrados"],
        "Valor": [len(df), df["empresa"].nunique(), df["cuenta"].nunique(), df["mes"].nunique()]
    })

    # ============================================================
    # FILTROS DE PANTALLA
    # ============================================================
    st.sidebar.header("🔍 Filtros de Análisis")

    empresas_opt = sorted(df_analyzed["empresa"].unique())
    selected_empresas = st.sidebar.multiselect("Casa / Empresa", empresas_opt, default=empresas_opt)

    vistas_opt = sorted(df_analyzed["vista"].unique())
    selected_vistas = st.sidebar.multiselect("Vista", vistas_opt, default=vistas_opt)

    segmentos_opt = sorted(df_analyzed["segmento"].unique())
    selected_segmentos = st.sidebar.multiselect("Segmento", segmentos_opt, default=segmentos_opt)

    alertas_opt = ["🔴 CRÍTICA", "🟠 IMPORTANTE", "🟡 MODERADA", "🟢 NORMAL"]
    selected_alertas = st.sidebar.multiselect("Nivel de Alerta", alertas_opt, default=alertas_opt)

    # Filtrar DataFrame principal
    filtered = df_analyzed[
        (df_analyzed["empresa"].isin(selected_empresas))
        & (df_analyzed["vista"].isin(selected_vistas))
        & (df_analyzed["segmento"].isin(selected_segmentos))
        & (df_analyzed["nivel_alerta"].isin(selected_alertas))
    ].copy()

    detail = df.copy()

    # ============================================================
    # TABS DE INTERFAZ
    # ============================================================
    tab1, tab2, tab3, tab4 = st.tabs(
        ["📈 Resumen por Casas", "🚨 Alertas & Variaciones", "🔍 Búsqueda / Detalle", "⚙️ Diagnóstico"]
    )

    with tab1:
        st.subheader("Visión General de Gastos por Casas y Logística")

        if not filtered.empty:
            pivote_pantalla = filtered.pivot_table(
                index=["segmento", "cuenta"],
                columns=["empresa", "vista"],
                values="monto",
                aggfunc="sum",
                fill_value=0,
            )
            st.dataframe(pivote_pantalla.style.format("{:,.2f}"), use_container_width=True)

            # Gráfico comparativo
            fig = px.bar(
                filtered,
                x="segmento",
                y="monto",
                color="empresa",
                barmode="group",
                facet_col="vista",
                title="Distribución de Gastos por Segmento, Casa y Vista",
            )
            st.plotly_chart(fig, use_container_width=True)
        else:
            st.warning("No hay datos disponibles con los filtros seleccionados.")

    with tab2:
        st.subheader("Alertas y Variaciones Significativas")
        if not filtered.empty:
            alert_df = filtered[filtered["nivel_alerta"].isin(["🔴 CRÍTICA", "🟠 IMPORTANTE"])]
            st.dataframe(
                alert_df[
                    [
                        "nivel_alerta",
                        "empresa",
                        "vista",
                        "segmento",
                        "cuenta",
                        "mes",
                        "monto",
                        "var_abs_mes_anterior",
                        "var_pct_mes_anterior",
                    ]
                ].style.format(
                    {
                        "monto": "{:,.2f}",
                        "var_abs_mes_anterior": "{:,.2f}",
                        "var_pct_mes_anterior": "{:.1%}",
                    }
                ),
                use_container_width=True,
            )
        else:
            st.info("Sin alertas para los filtros actuales.")

    with tab3:
        st.subheader("Exploración Detallada")
        st.dataframe(filtered, use_container_width=True)

    with tab4:
        st.subheader("Diagnóstico de Carga")
        st.table(diagnostics)

    # ============================================================
    # EXPORTACIÓN A EXCEL PROFESIONAL POR CASAS Y LOGÍSTICA
    # ============================================================
    st.divider()
    st.subheader("📥 Exportar análisis")

    if not filtered.empty:
        out = io.BytesIO()

        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            # 1. Matriz Resumen por Casas (Febeca, Sillaca, Beval con/sin Logística)
            pivote_casas = filtered.pivot_table(
                index=["segmento", "cuenta"],
                columns=["empresa", "vista"],
                values="monto",
                aggfunc="sum",
                fill_value=0,
            )
            pivote_casas.to_excel(writer, sheet_name="Resumen_Por_Casas")

            # 2. Hojas de Detalle del Análisis
            filtered.to_excel(writer, sheet_name="Variaciones", index=False)

            export_alerts = filtered[
                filtered["nivel_alerta"].isin(["🔴 CRÍTICA", "🟠 IMPORTANTE"])
            ].copy()
            export_alerts.to_excel(writer, sheet_name="Alertas", index=False)

            latest_export = (
                filtered[filtered["mes"] == filtered["mes"].max()]
                .sort_values("impacto_abs", ascending=False)
                .head(30)
            )
            latest_export.to_excel(writer, sheet_name="Top30_Ultimo_Mes", index=False)
            diagnostics.to_excel(writer, sheet_name="Diagnostico", index=False)

            if detail is not None and not detail.empty and not export_alerts.empty:
                parts = []
                for _, alert in export_alerts.head(100).iterrows():
                    rows = detail_matches_alert(alert, detail)
                    if not rows.empty:
                        parts.append(rows)
                if parts:
                    pd.concat(parts, ignore_index=True).drop_duplicates().to_excel(
                        writer, sheet_name="Detalle_Alertas", index=False
                    )

            # 3. Aplicación de Estilos y Formato Contable
            wb = writer.book

            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            num_fmt = '#,##0.00;(#,##0.00);"-"'
            pct_fmt = "0.0%"
            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9"),
            )

            for ws in wb.worksheets:
                ws.views.sheetView[0].showGridLines = True
                ws.freeze_panes = "C3" if ws.title == "Resumen_Por_Casas" else "A2"

                # Formatear encabezados de la primera fila/filas de título
                max_col = ws.max_column
                max_row_header = 2 if ws.title == "Resumen_Por_Casas" else 1

                for r in range(1, max_row_header + 1):
                    for col in range(1, max_col + 1):
                        cell = ws.cell(row=r, column=col)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Formatear celdas con números y autofit de columnas
                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = 0
                    for cell in col:
                        val = cell.value
                        if val is not None:
                            max_len = max(max_len, len(str(val)))

                        cell.border = thin_border

                        if isinstance(val, (int, float)):
                            header_name = str(ws.cell(row=1, column=cell.column).value or "").lower()
                            if "pct" in header_name or "%" in header_name:
                                cell.number_format = pct_fmt
                            else:
                                cell.number_format = num_fmt

                    ws.column_dimensions[col_letter].width = min(max(12, max_len + 3), 40)

            # Colorear alertas en la pestaña Alertas
            if "Alertas" in wb.sheetnames:
                ws_alt = wb["Alertas"]
                red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                red_font = Font(color="9C0006", bold=True)
                orange_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                orange_font = Font(color="9C6500", bold=True)

                for row in range(2, ws_alt.max_row + 1):
                    val = str(ws_alt.cell(row=row, column=1).value or "")
                    if "CRÍTICA" in val:
                        for col in range(1, ws_alt.max_column + 1):
                            ws_alt.cell(row=row, column=col).fill = red_fill
                            ws_alt.cell(row=row, column=col).font = red_font
                    elif "IMPORTANTE" in val:
                        for col in range(1, ws_alt.max_column + 1):
                            ws_alt.cell(row=row, column=col).fill = orange_fill
                            ws_alt.cell(row=row, column=col).font = orange_font

        out.seek(0)
        st.download_button(
            "📥 Descargar reporte Excel (Estructurado por Casas)",
            data=out.getvalue(),
            file_name=f"Analisis_Variaciones_Gastos_{datetime.now():%Y%m%d_%H%M}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
    else:
        st.info("No hay datos para exportar con los filtros actuales.")


if __name__ == "__main__":
    modulo_analizador_gastos()
